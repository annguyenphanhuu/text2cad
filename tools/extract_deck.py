"""One-off: turn the client's 2511 test deck into the HTML test bank tests/deck2511.html.

    python tools/extract_deck.py [--deck 2511_ATN_TEST_REPORT_Roman_EN.pptx] [--out tests/deck2511.html]

The bank is one self-contained page (see tests/harness/bank.py): every case with its prompt
(English, and the French original when tests/fixtures/pptx_*_cases.json has it), the drawing
the client's engineer made of the part (ground truth, embedded), the client's 2511 verdict
kept for reference, and room for the latest run result.  After this one run the .pptx is not
needed any more; tests/rerun_deck_ir.py reads the cases from the bank and writes results back.

Deck layout this relies on: per case a prompt box ending in "(OK)" / "(NOT OK)", a small
render of the old chatbot's output above a number label, and below the label the client's
drawing; two cases side by side on most slides (left / right half), one on some.  Cases whose
tester note says "complex shape" are listed under "excluded" and get no record: the pipeline
does not model those on purpose.
"""
import argparse
import json
import re
import sys
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from tests.harness import bank  # noqa: E402

SEC_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)[.\s]+([A-Z][A-Z \-/&']{2,})\s*$")
START_RE = re.compile(r"^\s*(\d+)\s*[.)]\s+(\S.*)$", re.S)
STATUS_RE = re.compile(r"\(\s*(OK|Not OK|NOT OK|NOK|KO)\s*\)\s*$", re.I)
LABEL_RE = re.compile(r"^(\d{1,3})([a-z])?$")      # deck case number, sometimes 167b for a variant
CHROME_TOP = 6.4          # logos / footer pictures sit below this line (inches)


def norm(t):
    return re.sub(r"\s+", " ", t.replace(" ", " ").replace("\x0b", " ")).strip()


def parse_cases(paras):
    """Cases inside one text frame: 'N. prompt ... (OK)' possibly spanning several paragraphs."""
    out, i = [], 0
    while i < len(paras):
        m = START_RE.match(paras[i])
        if not m or SEC_RE.match(paras[i]):
            i += 1
            continue
        num, body, status, j = int(m.group(1)), [m.group(2)], None, i + 1
        sm = STATUS_RE.search(paras[i])
        if sm:
            status = sm.group(1)
        while status is None and j < len(paras) and j - i <= 25:
            if START_RE.match(paras[j]) and STATUS_RE.search(paras[j]):
                break
            body.append(paras[j])
            sm = STATUS_RE.search(paras[j])
            if sm:
                status = sm.group(1)
            j += 1
        text = STATUS_RE.sub("", " ".join(body)).strip()
        out.append({"num": num, "prompt": norm(text),
                    "verdict": None if status is None else ("OK" if status.upper() == "OK" else "NOT OK")})
        i = j if status else i + 1
    return out


def read_slide(slide):
    """cases (with frame x), number labels, notes, content pictures - positions in inches."""
    cases, labels, notes, pics = [], [], [], []
    for sh in slide.shapes:
        x, y, w, h = (Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height))
        if sh.shape_type == 13:
            if y < CHROME_TOP:
                pics.append({"shape": sh, "x": x, "y": y, "cx": x + w / 2, "area": w * h, "px": sh.image.size[0]})
            continue
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        t = sh.text_frame.text.strip()
        if x > 12 or y > 6.8 or y < 0.5:
            continue                              # slide number, date, section title
        lm = LABEL_RE.match(t)
        if lm:
            labels.append({"n": "%03d%s" % (int(lm.group(1)), lm.group(2) or ""), "x": x, "y": y, "cx": x + w / 2})
            continue
        paras = [norm("".join(r.text for r in p.runs)) for p in sh.text_frame.paragraphs]
        paras = [p for p in paras if p]
        found = parse_cases(paras)
        if found:
            for k, c in enumerate(found):
                c["cx"] = x + w / 2 + 0.01 * k     # keep frame order when a frame holds two cases
                cases.append(c)
        else:
            notes.append({"text": norm(t), "cx": x + w / 2})
    if not cases and labels:
        # a few slides carry the prompt without the "N." prefix: the long text box is the prompt
        for n in [n for n in notes if len(n["text"]) > 60]:
            sm = STATUS_RE.search(n["text"])
            cases.append({"num": None, "prompt": STATUS_RE.sub("", n["text"]).strip(), "cx": n["cx"],
                          "verdict": None if sm is None else ("OK" if sm.group(1).upper() == "OK" else "NOT OK")})
            notes.remove(n)
    return cases, labels, notes, pics


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(ROOT / "2511_ATN_TEST_REPORT_Roman_EN.pptx"))
    ap.add_argument("--out", default=str(bank.BANK))
    args = ap.parse_args()

    fr = {}
    for name in ("pptx_ok_cases.json", "pptx_notok_cases.json"):
        p = ROOT / "tests" / "fixtures" / name
        if p.exists():
            for c in json.loads(p.read_text(encoding="utf-8"))["cases"]:
                fr[(c["slide"], c["num"])] = c["prompt"]
    august = {}
    sel = ROOT / "outputs" / "test_rerun_20260818" / "selection.json"
    if sel.exists():
        for c in json.loads(sel.read_text(encoding="utf-8"))["run"]:
            august[(c["slide"], c["num"])] = c["uid"]

    prs = Presentation(args.deck)
    half = Emu(prs.slide_width).inches / 2
    section, records, excluded, warnings = None, [], [], []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    m = SEC_RE.match(norm("".join(r.text for r in p.runs)))
                    if m:
                        section = "%s %s" % (m.group(1), m.group(2).strip())
        cases, labels, notes, pics = read_slide(slide)
        if not cases:
            continue
        cases.sort(key=lambda c: c["cx"])
        for k, c in enumerate(cases):
            if len(cases) == 1:
                mine = lambda e: True
            elif len(cases) == 2:
                mine = lambda e, col=k: (0 if e["cx"] < half else 1) == col
            else:
                warnings.append("slide %d: %d cases on one slide, pictures not assigned" % (si, len(cases)))
                mine = lambda e: False
            my_labels = sorted((l for l in labels if mine(l)), key=lambda l: l["y"])
            my_pics = [p for p in pics if mine(p)]
            note = " | ".join(n["text"] for n in notes if mine(n)) or None
            label = my_labels[0]["n"] if my_labels else None
            gt = None
            if my_pics:
                if my_labels:
                    below = [p for p in my_pics if p["y"] > my_labels[0]["y"]]      # the drawing sits under the number
                    gt = max(below, key=lambda p: p["area"]) if below else None
                else:
                    gt = max(my_pics, key=lambda p: p["px"])                        # the drawing is the wider image
            uid = label if label is not None else "s%03dn%s" % (si, c["num"] if c["num"] is not None else k + 1)
            rec = {"uid": uid, "slide": si, "section": section, "chapter": int(section.split()[0].split(".")[0]) if section else None,
                   "num": c["num"], "prompt_en": c["prompt"], "prompt_fr": fr.get((si, c["num"])),
                   "verdict_2511": c["verdict"], "tester_note": note, "august_uid": august.get((si, c["num"]))}
            if note and "complex" in note.lower():
                excluded.append(rec)
                continue
            if gt is None:
                warnings.append("slide %d case %s (uid %s): no ground-truth drawing found" % (si, c["num"], uid))
            rec["gt"] = bank.webp_uri(gt["shape"].image.blob, quality=80) if gt else None
            rec["result"], rec["history"] = None, []
            if c["verdict"] is None:
                warnings.append("slide %d case %s (uid %s): no (OK)/(NOT OK) mark" % (si, c["num"], uid))
            records.append(rec)

    uids = [r["uid"] for r in records]
    dup = sorted({u for u in uids if uids.count(u) > 1})
    if dup:
        warnings.append("duplicate uids: %s" % dup)
    data = {"title": "Deck 2511", "source": Path(args.deck).name, "created": time.strftime("%Y-%m-%d"),
            "cases": records, "excluded": excluded, "runs": []}
    size = bank.save(data, Path(args.out))
    print("%d cases (%d with a GT drawing, %d with a French prompt, %d in the August selection) | %d excluded as complex shapes | %s (%.1f MB)"
          % (len(records), sum(1 for r in records if r["gt"]), sum(1 for r in records if r["prompt_fr"]),
             sum(1 for r in records if r["august_uid"]), len(excluded), args.out, size / 1e6))
    for w in warnings:
        print("  !", w)


if __name__ == "__main__":
    main()
